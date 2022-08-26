import string

import openpyxl
from django.db.models import QuerySet, Count, Sum
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from basic import models as basic_models
from event import models as event_models
from event.file_generator.generators import helper
from event.file_generator.generators.abstract_generator import AbstractGenerator
from event.file_generator.generators.helper import get_participants_by_registration
from event.file_generator.models import FileTemplate


class AttributeGenerator(AbstractGenerator):

    def generate(self) -> Workbook:
        event: event_models.Event = self.generated_file.event
        registrations: QuerySet[event_models.Registration] = helper.get_registrations(event)
        file: FileTemplate = self.generated_file.template
        wb: Workbook = load_workbook(file.file)

        original: Worksheet = wb.active

        alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        font_style = Font(size="10", name='Calibri')

        booking_options: QuerySet[event_models.BookingOption] = event_models.BookingOption.objects.filter(event=event)
        booking_options_len = booking_options.count()

        original['A2'] = 'Index'
        original['B2'] = 'Stamm'
        original['C2'] = 'Teilnehmerzahl'
        original['D2'] = 'Anreise Art'
        original['E2'] = 'Anreise Zeit'
        original['F2'] = 'Kothen'
        original['G2'] = 'Jurten'
        index = 7
        for option in booking_options:
            original[f'{string.ascii_uppercase[index]}2'] = option.name
            index += 1

        participant_count = registrations.aggregate(count=Count('registrationparticipant'))['count'] or 0

        original['C1'] = f'{event.name} \n{helper.get_event_location(event)}'
        original['E1'] = helper.get_event_date(event)
        original['G1'] = participant_count

        index: int = 3
        registration: event_models.Registration
        for registration in registrations.all().order_by('scout_organisation__name'):
            travel_tag: basic_models.TravelAttribute = registration.tags \
                .instance_of(basic_models.TravelAttribute).first()
            tent_small_tag: basic_models.IntegerAttribute = registration.tags.filter(template_id=17).first()
            tent_large_tag: basic_models.IntegerAttribute = registration.tags.filter(template_id=19).first()

            all_participants = get_participants_by_registration(registration)
            original[f'A{index}'] = index - 2
            original[f'B{index}'] = helper.get_registration_scout_organistation_name(registration)
            original[f'C{index}'] = all_participants.count()

            for idx, option in enumerate(booking_options):
                option_count = all_participants.filter(booking_option=option).count()
                original[f'{string.ascii_uppercase[7 + idx]}{index}'] = option_count

            if travel_tag:
                original[f'D{index}'] = travel_tag.get_type_field_display()
                original[f'E{index}'] = travel_tag.get_time_field_display()

            if tent_small_tag:
                original[f'F{index}'] = tent_small_tag.integer_field

            if tent_large_tag:
                original[f'G{index}'] = tent_large_tag.integer_field

            original.row_dimensions[index].height = 25
            for letter in string.ascii_uppercase[:7 + booking_options_len]:
                original[f'{letter}{index}'].alignment = alignment
                original[f'{letter}{index}'].font = font_style

            index += 1

        mapper_ids = registrations.values_list('tags', flat=True)
        attributes = basic_models.AbstractAttribute.objects.filter(id__in=mapper_ids)

        count_tent_small: int = \
            attributes.filter(template_id=17).aggregate(count=Sum('integerattribute__integer_field'))['count'] or 0
        count_tent_large: int = \
            attributes.filter(template_id=19).aggregate(count=Sum('integerattribute__integer_field'))['count'] or 0

        original[f'A{index}'] = 'Ergebnis'
        original[f'C{index}'] = participant_count
        original[f'F{index}'] = count_tent_small
        original[f'G{index}'] = count_tent_large

        original.row_dimensions[index].height = 25
        for letter in string.ascii_uppercase[:7 + booking_options_len]:
            original[f'{letter}{index}'].alignment = alignment
            original[f'{letter}{index}'].font = font_style

        ref = f'A2:{string.ascii_uppercase[6 + booking_options_len]}{registrations.count() + 3}'
        medium_style = openpyxl.worksheet.table.TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        table = openpyxl.worksheet.table.Table(ref=ref, displayName='Attribute', tableStyleInfo=medium_style)
        original.add_table(table)

        return wb
