import string

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font

from event import models as event_models
from event.file_generator.generators import helper
from event.file_generator.generators.abstract_generator import AbstractGenerator
from event.file_generator.generators.helper import get_participants_by_event
from event.file_generator.models import FileTemplate


class ParticipantGenerator(AbstractGenerator):

    def generate(self) -> Workbook:
        event: event_models.Event = self.generated_file.event
        participants = get_participants_by_event(event)
        file: FileTemplate = self.generated_file.template
        wb: Workbook = load_workbook(file.file)

        original = wb.active
        original.sheet_format.defaultRowHeight = 25.0
        original.set_printer_settings(paper_size=original.PAPERSIZE_A4, orientation='landscape')

        table = original.tables['Teilnehmer']
        table.ref = f'A3:J{participants.count() + 3}'

        original['C1'] = f'{event.name} \n{helper.get_event_location(event)}'
        original['G1'] = helper.get_event_date(event)
        original['J1'] = participants.count()

        alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        font_style = Font(size="10", name='Calibri')

        participant: event_models.RegistrationParticipant
        for index, participant in enumerate(participants.all()):
            original[f'A{index + 4}'] = helper.get_participant_first_name(participant)
            original[f'B{index + 4}'] = helper.get_participant_last_name(participant)
            original[f'C{index + 4}'] = helper.get_participant_scout_name(participant)
            original[f'D{index + 4}'] = helper.get_participant_age(event, participant)
            original[f'E{index + 4}'] = helper.get_participant_gender(participant)
            original[f'F{index + 4}'] = helper.get_participant_mobile_numer(participant)
            original[f'G{index + 4}'] = helper.get_participant_email(participant)
            original[f'H{index + 4}'] = helper.get_participant_registration_scout_organisation_name(participant)
            original[f'I{index + 4}'] = helper.get_participant_street(participant)
            original[f'J{index + 4}'] = helper.get_participant_zip_city(participant)

            original.row_dimensions[index + 4].height = 25
            for letter in string.ascii_uppercase[:10]:
                original[f'{letter}{index + 4}'].alignment = alignment
                original[f'{letter}{index + 4}'].font = font_style

        return wb
